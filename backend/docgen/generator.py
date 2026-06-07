"""
EcoMind 多格式文书生成器

对标 Trae Solo 的 report-generator + slides + doc-coauthoring:
  生成格式: Markdown / HTML / PDF (via weasyprint) / Word (via python-docx)
  文种类型: 执法文书 / 环评报告 / 环境监测日报 / 会议纪要 / 政策解读

Usage:
    gen = DocGenerator()
    doc = await gen.generate("enforcement_case", data)
"""
from __future__ import annotations

import json
import logging
import os
from datetime import datetime
from pathlib import Path
from typing import Any, Optional

logger = logging.getLogger(__name__)

OUTPUT_DIR = Path(__file__).resolve().parent.parent / "outputs" / "documents"


class DocGenerator:
    """多格式文书生成器"""

    # ── 模板注册 ──

    TEMPLATES = {
        "enforcement_case": {
            "title": "环境执法案件文书",
            "sections": [
                "案件基本信息", "违法事实", "证据清单",
                "法律依据", "处罚决定", "告知事项",
            ],
        },
        "eia_report": {
            "title": "环境影响评价报告",
            "sections": [
                "项目概况", "工程分析", "环境现状",
                "环境影响预测", "环境保护措施", "结论",
            ],
        },
        "monitoring_daily": {
            "title": "环境监测日报",
            "sections": [
                "监测概况", "城市排名", "首要污染物",
                "AQI分布", "对比分析", "预报",
            ],
        },
        "inspection_record": {
            "title": "现场检查记录",
            "sections": [
                "被检查单位", "检查时间地点", "现场情况",
                "存在问题", "处理意见", "签字",
            ],
        },
        "meeting_minutes": {
            "title": "会议纪要",
            "sections": [
                "会议信息", "参会人员", "主要议题",
                "讨论要点", "决议事项", "下一步工作",
            ],
        },
        "policy_brief": {
            "title": "政策解读",
            "sections": [
                "政策背景", "主要内容", "重点条款",
                "实施要求", "时间安排", "咨询方式",
            ],
        },
    }

    def __init__(self):
        os.makedirs(OUTPUT_DIR, exist_ok=True)

    async def generate(
        self,
        doc_type: str,
        data: dict[str, Any],
        format: str = "markdown",
    ) -> dict[str, str]:
        """
        生成文书。

        Args:
            doc_type: enforcement_case / eia_report / monitoring_daily / inspection_record
            data: 填充数据 {title, date, location, content, ...}
            format: markdown / html / pdf / docx

        Returns:
            {"path": ..., "content": ..., "format": ...}
        """
        template = self.TEMPLATES.get(doc_type)
        if not template:
            return {"error": f"未知文种: {doc_type}。支持: {', '.join(self.TEMPLATES.keys())}"}

        title = data.get("title", template["title"])
        doc_date = data.get("date", datetime.now().strftime("%Y年%m月%d日"))

        if format == "markdown":
            content = self._render_markdown(title, doc_date, template, data)
            path = self._save(title, content, ".md")
            return {"path": str(path), "content": content, "format": "markdown",
                    "preview": content[:500]}

        elif format == "html":
            md_content = self._render_markdown(title, doc_date, template, data)
            html_content = self._md_to_html(md_content, title)
            path = self._save(title, html_content, ".html")
            return {"path": str(path), "content": html_content, "format": "html",
                    "preview": html_content[:500]}

        elif format == "pdf":
            md_content = self._render_markdown(title, doc_date, template, data)
            html_content = self._md_to_html(md_content, title)
            path = OUTPUT_DIR / f"{self._slug(title)}.pdf"
            try:
                self._html_to_pdf(html_content, path)
                return {"path": str(path), "content": "", "format": "pdf",
                        "preview": md_content[:300]}
            except ImportError:
                return {"error": "PDF 生成需要 weasyprint 库。请 pip install weasyprint",
                        "fallback_md": md_content[:500]}

        elif format == "docx":
            path = OUTPUT_DIR / f"{self._slug(title)}.docx"
            try:
                md_content = self._render_markdown(title, doc_date, template, data)
                self._md_to_docx(md_content, title, path)
                return {"path": str(path), "content": "", "format": "docx",
                        "preview": md_content[:300]}
            except ImportError:
                return {"error": "DOCX 生成需要 python-docx 库。请 pip install python-docx",
                        "fallback_md": md_content if isinstance(md_content, str) else ""}

        return {"error": f"不支持的格式: {format}"}

    def _render_markdown(self, title: str, date: str, template: dict, data: dict) -> str:
        """渲染 Markdown 文档"""
        lines = [f"# {title}", ""]
        lines.append(f"**日期**: {date}")
        if data.get("department"):
            lines.append(f"**部门**: {data['department']}")
        if data.get("case_id"):
            lines.append(f"**编号**: {data['case_id']}")
        lines.append("---\n")

        for section in template["sections"]:
            lines.append(f"## {section}")
            content_key = self._section_key(section)
            section_content = data.get(content_key, data.get("content", ""))
            if isinstance(section_content, list):
                for item in section_content:
                    lines.append(f"- {item}")
            elif isinstance(section_content, dict):
                for k, v in section_content.items():
                    lines.append(f"- **{k}**: {v}")
            else:
                lines.append(str(section_content) if section_content else "(待填写)")
            lines.append("")

        # 底部
        lines.append("---")
        lines.append(f"*本文档由 EcoMind OS 自动生成 | {datetime.now().isoformat()[:19]}*")
        return "\n".join(lines)

    def _section_key(self, section_name: str) -> str:
        """将中文节名转为 data dict 的 key"""
        mapping = {
            "案件基本信息": "basic_info",
            "违法事实": "violation_facts",
            "证据清单": "evidence_list",
            "法律依据": "legal_basis",
            "处罚决定": "penalty_decision",
            "告知事项": "notification",
            "项目概况": "project_overview",
            "工程分析": "engineering_analysis",
            "环境现状": "environmental_status",
            "环境影响预测": "impact_prediction",
            "环境保护措施": "protection_measures",
            "结论": "conclusion",
            "监测概况": "monitoring_overview",
            "城市排名": "city_ranking",
            "首要污染物": "primary_pollutant",
            "AQI分布": "aqi_distribution",
            "对比分析": "comparison_analysis",
            "预报": "forecast",
            "被检查单位": "inspected_unit",
            "检查时间地点": "inspection_info",
            "现场情况": "on_site_situation",
            "存在问题": "issues_found",
            "处理意见": "handling_opinion",
            "签字": "signature",
            "会议信息": "meeting_info",
            "参会人员": "attendees",
            "主要议题": "agenda",
            "讨论要点": "discussion_points",
            "决议事项": "resolutions",
            "下一步工作": "next_steps",
            "政策背景": "policy_background",
            "主要内容": "main_content",
            "重点条款": "key_clauses",
            "实施要求": "implementation_requirements",
            "时间安排": "timeline",
            "咨询方式": "contact_info",
        }
        return mapping.get(section_name, "content")

    def _md_to_html(self, md: str, title: str) -> str:
        """Markdown → HTML 用于 PDF 基础样式"""
        try:
            from markdown import markdown as md_lib
            body = md_lib(md, extensions=["tables", "fenced_code"])
        except ImportError:
            body = md.replace("\n", "<br>\n")
        return f"""<!DOCTYPE html><html><head><meta charset="utf-8">
<title>{title}</title>
<style>body{{font-family:'SimSun',serif;max-width:800px;margin:40px auto;line-height:2}}
h1{{text-align:center}} h2{{border-bottom:1px solid #999;padding-bottom:4px}}</style>
</head><body>{body}</body></html>"""

    def _html_to_pdf(self, html: str, path: Path):
        from weasyprint import HTML
        HTML(string=html).write_pdf(path)

    def _md_to_docx(self, md: str, title: str, path: Path):
        from docx import Document
        doc = Document()
        doc.add_heading(title, 0)
        for line in md.split("\n"):
            line = line.strip()
            if not line:
                continue
            if line.startswith("## "):
                doc.add_heading(line[3:], level=2)
            elif line.startswith("# "):
                doc.add_heading(line[2:], level=1)
            elif line.startswith("- "):
                doc.add_paragraph(line[2:], style="List Bullet")
            else:
                doc.add_paragraph(line)
        doc.save(path)

    def _save(self, title: str, content: str, ext: str) -> Path:
        path = OUTPUT_DIR / f"{self._slug(title)}{ext}"
        with open(path, "w", encoding="utf-8") as f:
            f.write(content)
        return path

    def _slug(self, title: str) -> str:
        import re
        slug = re.sub(r'[^\w一-鿿-]', '_', title)[:60]
        slug = re.sub(r'_+', '_', slug).strip('_')
        ts = datetime.now().strftime("%Y%m%d_%H%M")
        return f"{slug}_{ts}"

    def list_templates(self) -> list[dict]:
        return [
            {"type": k, "title": v["title"], "sections": len(v["sections"])}
            for k, v in self.TEMPLATES.items()
        ]
